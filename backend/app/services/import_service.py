from __future__ import annotations

import csv
import io
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.telegram.formatter import TreasuryFormatter


@dataclass(frozen=True)
class ImportRowPreview:
    row_number: int
    created_at: str
    tx_type: str
    amount: str
    category: str
    description: str | None
    source: str

    def to_dict(self) -> dict[str, str | int | None]:
        return asdict(self)


@dataclass(frozen=True)
class ImportPreview:
    filename: str
    detected_format: str
    supported: bool
    ready_to_import: bool
    total_rows: int
    importable_rows: list[ImportRowPreview]
    skipped_rows: list[str]
    warnings: list[str]
    rows_skipped_duplicates: int = 0
    rows_skipped_invalid: int = 0
    invalid_reasons: list[str] = field(default_factory=list)

    def to_pending_payload(self) -> list[dict[str, str | int | None]]:
        return [row.to_dict() for row in self.importable_rows]


VALIDATION_MESSAGES = {
    "missing_amount":   "нет суммы",
    "invalid_date":     "некорректная дата",
    "future_date":      "дата в будущем",
    "negative_amount":  "отрицательная сумма",
    "missing_desc":     "нет описания",
    "amount_too_large": "сумма > 10 000 000 ₽ (проверь)",
}


class ImportService:
    SUPPORTED_EXTENSIONS = {"csv", "xlsx", "pdf"}

    def __init__(self) -> None:
        self._formatter = TreasuryFormatter()
        self._reverse_categories = {
            value.lower(): key for key, value in self._formatter.CATEGORY_LABELS_RU.items()
        }

    @staticmethod
    def _validate_row(row: dict) -> list[str]:
        errors: list[str] = []
        amount_raw = row.get("amount") or row.get("Сумма")
        if not amount_raw:
            errors.append(VALIDATION_MESSAGES["missing_amount"])
        else:
            try:
                normalized = str(amount_raw).replace("₽", "").replace(" ", "").replace(",", ".")
                amount = float(normalized)
                if amount <= 0:
                    errors.append(VALIDATION_MESSAGES["negative_amount"])
                if amount > 10_000_000:
                    errors.append(VALIDATION_MESSAGES["amount_too_large"])
            except Exception:
                errors.append(VALIDATION_MESSAGES["missing_amount"])

        date_raw = row.get("date") or row.get("Дата")
        if not date_raw:
            errors.append(VALIDATION_MESSAGES["invalid_date"])
        else:
            try:
                raw = str(date_raw).strip()
                parsed: date | None = None
                try:
                    parsed = date.fromisoformat(raw)
                except ValueError:
                    for fmt in ("%d.%m.%Y", "%Y/%m/%d"):
                        try:
                            parsed = datetime.strptime(raw, fmt).date()
                            break
                        except ValueError:
                            continue
                if parsed is None:
                    errors.append(VALIDATION_MESSAGES["invalid_date"])
                elif parsed > date.today():
                    errors.append(VALIDATION_MESSAGES["future_date"])
            except Exception:
                errors.append(VALIDATION_MESSAGES["invalid_date"])

        desc_raw = row.get("description") or row.get("Описание")
        if not desc_raw or not str(desc_raw).strip():
            errors.append(VALIDATION_MESSAGES["missing_desc"])

        return errors

    @staticmethod
    def _dedup_key(created_at: str, amount: str, description: str | None) -> str:
        desc = (description or "").lower().strip()
        return f"{created_at}:{amount}:{desc}"

    def apply_dedup(self, preview: ImportPreview, existing_keys: set[str]) -> ImportPreview:
        """Return a new ImportPreview with duplicate rows removed from importable_rows."""
        if not existing_keys:
            return preview
        to_import: list[ImportRowPreview] = []
        dup_count = 0
        for row in preview.importable_rows:
            key = self._dedup_key(row.created_at, row.amount, row.description)
            if key in existing_keys:
                dup_count += 1
            else:
                to_import.append(row)
        if dup_count == 0:
            return preview
        return ImportPreview(
            filename=preview.filename,
            detected_format=preview.detected_format,
            supported=preview.supported,
            ready_to_import=bool(to_import),
            total_rows=preview.total_rows,
            importable_rows=to_import,
            skipped_rows=preview.skipped_rows,
            warnings=preview.warnings,
            rows_skipped_duplicates=dup_count,
            rows_skipped_invalid=preview.rows_skipped_invalid,
            invalid_reasons=list(preview.invalid_reasons),
        )

    @staticmethod
    def detect_format(filename: str | None) -> str:
        if not filename or "." not in filename:
            return "unknown"
        return filename.rsplit(".", 1)[-1].lower()

    def build_preview(self, *, filename: str, content: bytes) -> ImportPreview:
        detected_format = self.detect_format(filename)
        if detected_format == "pdf":
            return ImportPreview(
                filename=filename,
                detected_format=detected_format,
                supported=False,
                ready_to_import=False,
                total_rows=0,
                importable_rows=[],
                skipped_rows=[],
                warnings=[
                    "PDF-импорт пока в безопасном режиме: файл распознан, но автопарсинг отключён.",
                    "Сейчас можно импортировать только CSV/XLSX из понятной табличной структуры.",
                ],
            )
        if detected_format == "csv":
            return self._build_csv_preview(filename=filename, content=content)
        if detected_format == "xlsx":
            return self._build_xlsx_preview(filename=filename, content=content)
        return ImportPreview(
            filename=filename,
            detected_format=detected_format,
            supported=False,
            ready_to_import=False,
            total_rows=0,
            importable_rows=[],
            skipped_rows=[],
            warnings=["Поддерживаются только CSV, XLSX и PDF."],
        )

    def render_preview_message(self, preview: ImportPreview) -> str:
        if not preview.supported:
            lines = [f"📥 Импорт: {preview.filename}", "", *preview.warnings]
            return "\n".join(lines)

        lines = [
            f"📥 Dry-run импорт: {preview.filename}",
            f"Формат: {preview.detected_format}",
            f"Всего строк: {preview.total_rows}",
            f"Будет импортировано: {len(preview.importable_rows)}",
        ]
        if preview.rows_skipped_duplicates:
            lines.append(f"Пропущено дублей: {preview.rows_skipped_duplicates} (уже есть в базе)")
        if preview.rows_skipped_invalid:
            reasons_str = ", ".join(preview.invalid_reasons) if preview.invalid_reasons else "ошибки в данных"
            lines.append(f"Пропущено невалидных: {preview.rows_skipped_invalid} ({reasons_str})")
        if preview.skipped_rows and not preview.rows_skipped_invalid:
            lines.append(f"Пропущено: {len(preview.skipped_rows)}")
        if preview.importable_rows:
            lines.append("\nПредпросмотр:")
            for row in preview.importable_rows[:5]:
                kind = "доход" if row.tx_type == "income" else "расход"
                lines.append(
                    f"• {row.created_at} · {kind} · {row.amount} ₽ · {self._formatter.category_label(row.category)}"
                    + (f" · {row.description}" if row.description else "")
                )
            if len(preview.importable_rows) > 5:
                lines.append(f"…и ещё {len(preview.importable_rows) - 5}")
        if preview.skipped_rows:
            lines.append("\nЧто пропущено:")
            lines.extend(f"• {item}" for item in preview.skipped_rows[:5])
            if len(preview.skipped_rows) > 5:
                lines.append(f"…и ещё {len(preview.skipped_rows) - 5}")
        if preview.warnings:
            lines.append("\nЗамечания:")
            lines.extend(f"• {item}" for item in preview.warnings)
        if preview.ready_to_import:
            lines.append("\nЕсли всё ок, нажми подтверждение. До этого в базу ничего не пишется.")
        return "\n".join(lines)

    def _build_csv_preview(self, *, filename: str, content: bytes) -> ImportPreview:
        decoded = self._decode_csv(content)
        reader = csv.DictReader(io.StringIO(decoded))
        if not reader.fieldnames:
            return self._empty_preview(filename, "csv", warning="Не вижу заголовков колонок в CSV.")
        rows = list(reader)
        return self._preview_from_rows(filename=filename, detected_format="csv", rows=rows)

    def _build_xlsx_preview(self, *, filename: str, content: bytes) -> ImportPreview:
        wb = load_workbook(io.BytesIO(content), data_only=True)
        ws = wb["Транзакции"] if "Транзакции" in wb.sheetnames else wb[wb.sheetnames[0]]

        header_row_idx = None
        headers: list[str] | None = None
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [str(cell).strip() if cell is not None else "" for cell in row]
            if {"Дата", "Тип", "Сумма", "Категория"}.issubset(set(values)):
                header_row_idx = idx
                headers = values
                break
        if header_row_idx is None or headers is None:
            return self._empty_preview(filename, "xlsx", warning="Не нашёл строку заголовков в XLSX.")

        rows: list[dict[str, Any]] = []
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not any(cell not in (None, "") for cell in row):
                continue
            values = list(row) + [None] * (len(headers) - len(row))
            rows.append({headers[i]: values[i] for i in range(len(headers))})
        return self._preview_from_rows(filename=filename, detected_format="xlsx", rows=rows)

    def _preview_from_rows(self, *, filename: str, detected_format: str, rows: list[dict[str, Any]]) -> ImportPreview:
        importable_rows: list[ImportRowPreview] = []
        skipped_rows: list[str] = []
        warnings: list[str] = []
        rows_skipped_invalid = 0
        invalid_reasons: list[str] = []

        for idx, row in enumerate(rows, start=2):
            try:
                importable_rows.append(self._normalize_row(idx, row))
            except ValueError as exc:
                rows_skipped_invalid += 1
                # Collect unique human-readable reasons from _validate_row
                validation_errors = self._validate_row(row)
                if validation_errors:
                    for reason in validation_errors:
                        if reason not in invalid_reasons:
                            invalid_reasons.append(reason)
                skipped_rows.append(f"строка {idx}: {exc}")

        if not importable_rows:
            warnings.append("В файле не нашлось ни одной валидной записи для импорта.")
        return ImportPreview(
            filename=filename,
            detected_format=detected_format,
            supported=True,
            ready_to_import=bool(importable_rows),
            total_rows=len(rows),
            importable_rows=importable_rows,
            skipped_rows=skipped_rows,
            warnings=warnings,
            rows_skipped_invalid=rows_skipped_invalid,
            invalid_reasons=invalid_reasons,
        )

    def _normalize_row(self, row_number: int, row: dict[str, Any]) -> ImportRowPreview:
        date_value = self._string(row, "Дата")
        tx_type = self._normalize_type(self._string(row, "Тип"))
        amount = self._normalize_amount(self._string(row, "Сумма"))
        category = self._normalize_category(self._string(row, "Категория"))
        description = self._optional_string(row, "Описание")
        source = self._normalize_source(self._optional_string(row, "Источник") or "text")
        created_at = self._normalize_date(date_value)
        return ImportRowPreview(
            row_number=row_number,
            created_at=created_at.isoformat(),
            tx_type=tx_type,
            amount=f"{amount:.2f}",
            category=category,
            description=description,
            source=source,
        )

    @staticmethod
    def _string(row: dict[str, Any], key: str) -> str:
        value = row.get(key)
        if value is None:
            raise ValueError(f"нет колонки или значения «{key}»")
        text = str(value).strip()
        if not text:
            raise ValueError(f"пустое значение в колонке «{key}»")
        return text

    @staticmethod
    def _optional_string(row: dict[str, Any], key: str) -> str | None:
        value = row.get(key)
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _normalize_type(value: str) -> str:
        normalized = value.strip().lower()
        if normalized in {"income", "доход"}:
            return "income"
        if normalized in {"expense", "расход"}:
            return "expense"
        raise ValueError(f"непонятный тип операции «{value}»")

    @staticmethod
    def _normalize_amount(value: str) -> Decimal:
        normalized = value.replace("₽", "").replace(" ", "").replace(",", ".")
        try:
            amount = Decimal(normalized)
        except InvalidOperation as exc:
            raise ValueError(f"непонятная сумма «{value}»") from exc
        if amount <= 0:
            raise ValueError("сумма должна быть больше нуля")
        if amount > 10_000_000:
            raise ValueError(f"amount too large: {amount}")
        return amount

    def _normalize_category(self, value: str) -> str:
        normalized = value.strip()
        lower = normalized.lower()
        if lower in self._reverse_categories:
            return self._reverse_categories[lower]
        slug_like = lower.replace("-", "_").replace(" ", "_")
        return slug_like

    @staticmethod
    def _normalize_source(value: str) -> str:
        normalized = value.strip().lower()
        if normalized in {"voice", "голос", "voice_note"}:
            return "voice"
        # Import flow intentionally persists only enum-safe short sources.
        # "webapp" breaks current DB/source enum storage in production, so
        # anything non-voice is safely downgraded to text until schema is widened.
        return "text"

    @staticmethod
    def _normalize_date(value: str) -> date:
        raw = value.strip()
        parsed_date: date | None = None
        try:
            parsed_date = date.fromisoformat(raw)
        except ValueError:
            pass
        if parsed_date is None:
            for parser in ("%d.%m.%Y", "%Y/%m/%d"):
                try:
                    parsed_date = datetime.strptime(raw, parser).date()
                    break
                except ValueError:
                    continue
        if parsed_date is None:
            raise ValueError(f"непонятная дата «{value}»")
        if parsed_date > date.today():
            raise ValueError(f"future date: {parsed_date}")
        return parsed_date

    @staticmethod
    def _decode_csv(content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1251"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")

    @staticmethod
    def _empty_preview(filename: str, detected_format: str, *, warning: str) -> ImportPreview:
        return ImportPreview(
            filename=filename,
            detected_format=detected_format,
            supported=True,
            ready_to_import=False,
            total_rows=0,
            importable_rows=[],
            skipped_rows=[],
            warnings=[warning],
        )
