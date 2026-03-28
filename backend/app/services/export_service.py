from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.schemas.transaction import TransactionItem
from app.telegram.formatter import TreasuryFormatter


@dataclass(frozen=True)
class ExportArtifact:
    filename: str
    content: bytes
    mime_type: str
    caption: str


class ExportService:
    def __init__(self) -> None:
        self._formatter = TreasuryFormatter()

    def build(self, items: list[TransactionItem], *, period: str, export_format: str) -> ExportArtifact:
        export_format = self.normalize_format(export_format)
        if export_format == "xlsx":
            return self._build_xlsx(items, period=period)
        return self._build_csv(items, period=period)

    @staticmethod
    def normalize_format(value: str | None) -> str:
        value = (value or "csv").strip().lower()
        if value in {"xlsx", "excel"}:
            return "xlsx"
        return "csv"

    @staticmethod
    def normalize_period(value: str | None) -> str:
        value = (value or "month").strip().lower()
        return value if value in {"today", "week", "month", "all"} else "month"

    @staticmethod
    def period_label(period: str) -> str:
        return {"today": "сегодня", "week": "неделю", "month": "месяц", "all": "всё время"}.get(period, period)

    @staticmethod
    def period_slug(period: str) -> str:
        return {"today": "today", "week": "week", "month": "month", "all": "all"}.get(period, period)

    def _rows(self, items: list[TransactionItem]) -> list[list[str]]:
        rows: list[list[str]] = []
        for item in items:
            date_str = self._format_date(item.created_at)
            type_label = "Доход" if item.type == "income" else "Расход"
            category_label = self._formatter.category_label(item.category)
            amount_label = self._format_amount(item.amount)
            rows.append([
                date_str,
                type_label,
                amount_label,
                category_label,
                item.description or "",
                item.source,
            ])
        return rows

    def _build_csv(self, items: list[TransactionItem], *, period: str) -> ExportArtifact:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(["Дата", "Тип", "Сумма", "Категория", "Описание", "Источник"])
        writer.writerows(self._rows(items))
        csv_bytes = buf.getvalue().encode("utf-8-sig")
        return ExportArtifact(
            filename=f"finance_{self.period_slug(period)}.csv",
            content=csv_bytes,
            mime_type="text/csv",
            caption=f"📊 Экспорт за {self.period_label(period)}: {len(items)} записей",
        )

    def _build_xlsx(self, items: list[TransactionItem], *, period: str) -> ExportArtifact:
        wb = Workbook()
        ws = wb.active
        ws.title = "Транзакции"

        summary_rows = [
            ["Период", self.period_label(period)],
            ["Количество записей", len(items)],
            ["Экспортировано", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")],
        ]
        for row in summary_rows:
            ws.append(row)
        ws.append([])

        headers = ["Дата", "Тип", "Сумма", "Категория", "Описание", "Источник"]
        ws.append(headers)
        for cell in ws[5]:
            cell.font = Font(bold=True)

        for row in self._rows(items):
            ws.append(row)

        ws.freeze_panes = "A6"
        self._autosize_columns(ws)

        out = io.BytesIO()
        wb.save(out)
        return ExportArtifact(
            filename=f"finance_{self.period_slug(period)}.xlsx",
            content=out.getvalue(),
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            caption=f"📊 Excel-экспорт за {self.period_label(period)}: {len(items)} записей",
        )

    @staticmethod
    def _autosize_columns(ws) -> None:
        for idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 40)

    @staticmethod
    def _format_date(raw: str) -> str:
        return raw[:10] if raw else ""

    @staticmethod
    def _format_amount(raw: str) -> str:
        return f"{Decimal(raw):.2f}"
