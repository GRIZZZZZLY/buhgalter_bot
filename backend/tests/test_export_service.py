from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.schemas.transaction import TransactionItem
from app.services.export_service import ExportService


def _items() -> list[TransactionItem]:
    return [
        TransactionItem(
            id=uuid4(),
            telegram_user_id=555502880,
            type="expense",
            amount="1200.00",
            category="food",
            description="Самокат",
            source="text",
            raw_input="еда 1200",
            confidence=0.98,
            created_at="2026-03-28T10:11:12+00:00",
            updated_at="2026-03-28T10:11:12+00:00",
        ),
        TransactionItem(
            id=uuid4(),
            telegram_user_id=555502880,
            type="income",
            amount="5000.00",
            category="salary",
            description="Аванс",
            source="text",
            raw_input="аванс 5000",
            confidence=0.99,
            created_at="2026-03-28T11:11:12+00:00",
            updated_at="2026-03-28T11:11:12+00:00",
        ),
    ]


def test_export_service_builds_csv_artifact() -> None:
    artifact = ExportService().build(_items(), period="month", export_format="csv")

    assert artifact.filename == "finance_month.csv"
    assert artifact.mime_type == "text/csv"
    text = artifact.content.decode("utf-8-sig")
    assert "Дата,Тип,Сумма,Категория,Описание,Источник" in text
    assert "Расход,1200.00" in text
    assert "Доход,5000.00" in text


def test_export_service_builds_xlsx_artifact() -> None:
    artifact = ExportService().build(_items(), period="week", export_format="xlsx")

    assert artifact.filename == "finance_week.xlsx"
    assert artifact.mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(BytesIO(artifact.content))
    ws = wb["Транзакции"]
    assert ws["A1"].value == "Период"
    assert ws["B1"].value == "неделю"
    assert ws["A5"].value == "Дата"
    assert ws["B6"].value == "Расход"
    assert ws["B7"].value == "Доход"


def test_export_service_normalizes_excel_alias() -> None:
    assert ExportService.normalize_format("excel") == "xlsx"
    assert ExportService.normalize_period("all") == "all"
